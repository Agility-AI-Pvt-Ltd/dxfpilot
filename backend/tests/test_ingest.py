"""Workbook recognition without fixed header words."""

from __future__ import annotations

import io
import os

os.environ["OPENAI_API_KEY"] = ""
os.environ["LANGSMITH_TRACING"] = "false"

import openpyxl  # noqa: E402
import pytest  # noqa: E402

from cadpilot.ingest import excel  # noqa: E402
from cadpilot.ingest.excel import load_workbooks, parse_quantity  # noqa: E402

EQUIPMENT = [
    ("RECEPTION SECTION", None, None, None),
    ("Tanker Unloading Pump (VFD Operated)", "40 KLPH", 2, "Nos."),
    ("Duplex Inline Strainer (with Manual Changeover)", "40 KLPH", 2, "Nos."),
    ("Raw Milk Storage Silo with Bird Cage", "100 KL", 3, "Nos."),
    ("PROCESSING SECTION", None, None, None),
    ("Raw Milk Transfer Pump (from RMST to Pasteurizer)", "20 KLPH", 2, "Nos."),
    ("Milk Pasteurizer with all Accessories", "20 KLPH", 2, "Nos."),
]


def workbook(sheets: dict[str, list[tuple]]) -> tuple[str, bytes]:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return f"{next(iter(sheets))}.xlsx", buf.getvalue()


def equipment_sheet(header: tuple | None, order=(0, 1, 2, 3), rows=EQUIPMENT) -> list[tuple]:
    body = [tuple(r[i] for i in order) for r in rows]
    top = [("Plant equipment list",), ()]
    return top + ([tuple(header[i] for i in order)] if header else []) + body


def described(d) -> dict[str, tuple]:
    return {r.description: (str(r.capacity), r.qty) for r in d.equipment_rows}


@pytest.mark.parametrize(
    "header",
    [
        ("DESCRIPTION", "CAPACITY", "QTY.", "UOM"),
        ("Item Description", "Capacity", "Quantity", "Unit"),
        ("Particulars", "Rating", "Nos", "UoM"),
        ("Descripton", "Capacty", "Qty", "Unit"),  # typos
        ("विवरण", "क्षमता", "मात्रा", "इकाई"),  # Hindi
    ],
)
def test_equipment_headers_with_variations(header):
    d = load_workbooks([workbook({"Equipment": equipment_sheet(header)})], use_llm=False)
    assert described(d)["Raw Milk Storage Silo with Bird Cage"] == ("100 KL", 3.0)
    assert len(d.equipment_rows) == 5
    assert d.detections[0].method == "headers"
    assert d.equipment_rows[0].section == "RECEPTION SECTION"


def test_shuffled_columns_and_unit_in_header():
    rows = [(r[0], float(r[1].split()[0]) if r[1] else None, r[2], r[3]) for r in EQUIPMENT]  # bare numbers
    # columns in the order Qty, Description, Capacity, UOM; capacities are bare numbers
    sheet = equipment_sheet(("Description", "Capacity (KLPH)", "Qty", "UOM"), order=(2, 0, 1, 3), rows=rows)
    d = load_workbooks([workbook({"Equipment": sheet})], use_llm=False)
    assert described(d)["Milk Pasteurizer with all Accessories"] == ("20 KLPH", 2.0)
    assert d.detections[0].columns["capacity unit"] == "KLPH"


def test_table_without_any_header_is_read_from_its_content():
    d = load_workbooks([workbook({"Sheet1": equipment_sheet(None)})], use_llm=False)
    assert d.detections[0].method == "content"
    assert described(d)["Tanker Unloading Pump (VFD Operated)"] == ("40 KLPH", 2.0)


def test_civil_boq_without_capacity_is_not_taken_as_equipment():
    civil = [("DESCRIPTION", "QTY", "UNIT", "UNIT RATE"), ("Excavation in all soils", 1200, "cum", 450),
             ("Plain cement concrete M15", 300, "cum", 5200), ("Reinforcement steel Fe500", 45, "MT", 62000)]
    d = load_workbooks([workbook({"Civil": civil})], use_llm=False)
    assert not d.equipment_rows and "No equipment list recognised" in d.warnings[0]


def test_mass_balance_headers_in_hindi_with_percent_values():
    mb = [("दूध", "लीटर", "किलो", "वसा", "एसएनएफ"), ("Milk Reception", 500000, 515000, 4.0, 8.2)]
    d = load_workbooks([workbook({"MB": mb})], use_llm=False)
    assert d.daily_intake_litres() == 500000
    assert d.mass_balance_inputs[0].fat == pytest.approx(0.04) and d.mass_balance_inputs[0].snf == pytest.approx(0.082)


@pytest.mark.parametrize(
    ("raw", "expected"),
    [("40 KLPH", "40 KLPH"), ("40 kl/hr", "40 KLPH"), ("1,500 LPH", "1500 LPH"), ("100 Kilolitre", "100 KL"),
     ("200 Ltr.", "200 L"), ("40 m³/hr", "40 m3/h"), ("5 / 10 KLPH", "5 KLPH"), ("Suitable", "None"),
     ("8 TPH", "8 t/h"), ("540 CFM", "540 CFM"), ("1010 KVA", "1010 kVA"), ("30 MTPD", "30 t/day"), ("20 Tin/min", "20 per min")],
)
def test_capacity_units(raw, expected):
    assert str(parse_quantity(raw)) == expected


# ---- LLM mapping (simulated model; the real call has the same schema) --------------------------------------

FRENCH = ("Désignation", "Débit", "Quantité", "Unité")  # not in the synonym list


def fake_llm_mapping(monkeypatch, sheets_answer):
    monkeypatch.setattr(excel, "_llm_map", lambda sheets, need: excel._LLMWorkbookMap(sheets=sheets_answer))


def test_unknown_language_is_mapped_by_the_llm(monkeypatch):
    fake_llm_mapping(monkeypatch, [excel._LLMSheetMap(
        file="Equipements.xlsx", sheet="Equipements", kind="equipment_list", header_row=3, first_data_row=4,
        equipment=excel._LLMEquipmentColumns(description="A", capacity="B", quantity="C", unit="D", capacity_unit=None),
        mass_balance=[],
    )])
    d = load_workbooks([workbook({"Equipements": equipment_sheet(FRENCH)})])
    assert d.detections[0].method == "llm"
    assert d.detections[0].columns["description"] == "A: Désignation"
    assert described(d)["Milk Pasteurizer with all Accessories"] == ("20 KLPH", 2.0)


def test_wrong_llm_mapping_is_rejected(monkeypatch):
    """The model swapped description and quantity: the values contradict it, so it is not used."""
    fake_llm_mapping(monkeypatch, [excel._LLMSheetMap(
        file="Equipements.xlsx", sheet="Equipements", kind="equipment_list", header_row=3, first_data_row=4,
        equipment=excel._LLMEquipmentColumns(description="C", capacity="D", quantity="A", unit=None, capacity_unit=None),
        mass_balance=[],
    )])
    d = load_workbooks([workbook({"Equipements": equipment_sheet(FRENCH)})])
    # rejected → falls through to reading the content, which recovers the right columns
    assert d.detections[0].method == "content"
    assert described(d)["Raw Milk Storage Silo with Bird Cage"] == ("100 KL", 3.0)


def test_bundled_dummy_data_unchanged():
    from pathlib import Path

    d = load_workbooks(sorted(Path(__file__).parents[2].joinpath("data").glob("*.xlsx")), use_llm=False)
    # 295 rows with a quantity + 9 lump-sum rows (Automation Hardware, UPS, ...) the old reader silently dropped
    assert len(d.equipment_rows) == 304 and sum(1 for r in d.equipment_rows if r.qty is not None) == 295
    assert {r.ref.split("!")[1] for r in d.equipment_rows if r.qty is None} >= {"R255", "R260", "R277"}
    assert all(r.flags == ["quantity missing"] for r in d.equipment_rows if r.qty is None)
    assert d.daily_intake_litres() == 500000
    assert len(d.mass_balance_outputs) == 13
    assert d.warnings == [f"{len(d.flagged_rows())} equipment row(s) need attention: kept, but a value could not be read."]
    assert {t.kind for t in d.detections} == {"equipment_list", "mass_balance", "design_criteria"}


def test_llm_mapping_keeps_section_headings(monkeypatch):
    """Seen live: the model pointed first_data_row past a heading; reading starts below the header anyway."""
    fake_llm_mapping(monkeypatch, [excel._LLMSheetMap(
        file="Equipements.xlsx", sheet="Equipements", kind="equipment_list", header_row=3, first_data_row=5,
        equipment=excel._LLMEquipmentColumns(description="A", capacity="B", quantity="C", unit="D", capacity_unit=None),
        mass_balance=[],
    )])
    d = load_workbooks([workbook({"Equipements": equipment_sheet(FRENCH)})])
    assert d.equipment_rows[0].section == "RECEPTION SECTION"


# ---- ① never drop a row ------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "qty", "note"),
    [(2, 2, ""), ("3 Nos", 3, ""), ("2W+1S", 3, "2 working + 1 standby"), ("2 W + 1 F", 3, "2 working + 1 future"),
     ("2 (1W+1S)", 2, "1w+1s"), ("2+1", 3, "2 + 1"), ("two", 2, "")],
)
def test_quantity_expressions(raw, qty, note):
    assert excel.parse_count(raw) == (qty, note)


@pytest.mark.parametrize(
    ("raw", "cap", "note_part"),
    [("2 x 20 KLPH", "20 KLPH", "2 × 20 KLPH"), ("20 m3/h @ 3 bar", "20 m3/h", "at 3 bar"),
     ("15-20 KLPH", "20 KLPH", "range 15–20"), ("200 lts", "200 L", "")],
)
def test_capacity_expressions(raw, cap, note_part):
    q, note = excel.parse_capacity(raw)
    assert str(q) == cap and note_part in note


def test_unreadable_rows_are_kept_and_flagged():
    rows = EQUIPMENT + [
        ("Homogenizer", "20 KLPH", "as per vendor", "Nos."),  # quantity text
        ("Cream separator", "about twenty KLPH", 1, "Nos."),  # capacity text with no digits: kept, not flagged
        ("CIP return pump", "15 kl per shift", 2, "Nos."),  # capacity with digits but an unknown unit
        ("Automation hardware", "-----", None, None),  # lump sum, no quantity
    ]
    d = load_workbooks([workbook({"Equipment": equipment_sheet(("DESCRIPTION", "CAPACITY", "QTY", "UOM"), rows=rows)})], use_llm=False)
    by = {r.description: r for r in d.equipment_rows}
    assert len(d.equipment_rows) == 9  # 5 items + all 4 awkward rows; only the 2 headings are not items
    assert by["Homogenizer"].flags == ["quantity not understood: 'as per vendor'"]
    assert by["Cream separator"].flags == [] and by["Cream separator"].capacity is None
    assert by["CIP return pump"].flags == ["capacity not understood: '15 kl per shift'"]
    assert by["Automation hardware"].flags == ["quantity missing"]
    assert by["Raw Milk Storage Silo with Bird Cage"].section == "RECEPTION SECTION"
